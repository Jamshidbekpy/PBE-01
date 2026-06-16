from main_abc import ProductStructure
from openpyxl import Workbook, load_workbook


class Product(ProductStructure):
    def __init__(self, product_family, product_name, product_quantity, product_price, selling_price, product_sold, product_unit):
        self.product_family = product_family
        self.product_name = product_name
        self.product_quantity = product_quantity
        self.product_price = product_price
        self.selling_price = selling_price
        self.product_sold = product_sold
        self.product_unit = product_unit

    @classmethod
    def get_product_family(cls):
        return cls.product_family
    
    @classmethod
    def get_product_name(cls):
        return cls.product_name
    
    @classmethod
    def get_product_quantity(cls):
        return cls.product_quantity
    
    @classmethod
    def get_product_price(cls):
        return cls.product_price
    
    @classmethod
    def get_selling_price(cls):
        return cls.selling_price
    
    @classmethod
    def get_product_sold(cls):
        return cls.product_sold
    
    @classmethod
    def get_product_unit(cls):
        return cls.product_unit
    
    
    def add_product(self):
        wb = load_workbook('full_products.xlsx')
        ws = wb.active
        
        print(self.product_family in wb.sheetnames, "##################################")
        print(self.product_name in [row[0].value for row in wb[f'{self.product_family}'].iter_rows()], "##################################")
        
        
        if not(self.product_family in wb.sheetnames) and not(self.product_name in [row[0].value for row in wb[f'{self.product_family}'].iter_rows()]):
            wb.create_sheet(self.product_family)
            ws = wb[f'{self.product_family}']
            ws.append(['Name', 'Quantity', 'Buy Price', 'Sell Price', 'Sold', 'Unit'])
            ws.append([self.product_name, self.product_quantity, self.product_price, self.selling_price, self.product_sold, self.product_unit])
            wb.save('full_products.xlsx')
            
        elif self.product_family in wb.sheetnames and not(self.product_name in [row[0].value for row in wb[f'{self.product_family}'].iter_rows()]):
            ws = wb[f'{self.product_family}']
            ws.append([self.product_name, self.product_quantity, self.product_price, self.selling_price, self.product_sold, self.product_unit])
            wb.save('full_products.xlsx')
            
        else:
            row_index = None
            ws = wb[f'{self.product_family}']
            product_name = self.product_name
            for cell in ws['A']:
                if cell.value == product_name:
                    row_index = cell.row
                    break
            
            current_value = int(ws[f'B{row_index}'].value)

            ws[f'B{row_index}'] = current_value + self.product_quantity 
            
            wb.save("full_products.xlsx") 
            print(f"{product_name} miqdori {current_value+self.product_quantity} ga o'zgartirildi va saqlandi.")
            
            
            
        wb.close()
        
        
            



